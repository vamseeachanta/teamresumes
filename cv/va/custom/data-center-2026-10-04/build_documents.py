"""Refine verified candidate-authored resume material for three applications."""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT.parent / "anthropic-5375376008"
COMMON = (
    "Texas-licensed Professional Engineer with 23 years of marine and offshore engineering "
    "experience and 10+ years leading multidisciplinary teams from scope development through "
    "delivery. Combines recent CFD modelling and analysis with academic thesis work in fluid "
    "mechanics and heat transfer. Experience spans finite-element analysis, design reviews, "
    "emergency engineering response and analysis automation. "
)
CFD = (
    "My recent OpenFOAM work spans laminar, turbulent and free-surface flows, including "
    "benchmark-focused boundary-layer, vortex-shedding, airfoil and dam-break cases. My "
    "engineering work also includes COMSOL multiphysics modelling of fluid mechanics, mass "
    "transfer, diffusion and electrochemistry. Academic thesis work in fluid mechanics and "
    "heat transfer provides an additional foundation for this transition."
)
DELIVERY = (
    "At 2H Offshore, I progressed from analyst to engineering manager and technical authority, "
    "with responsibility for design reviews, technical reports, client communication and "
    "delivery across more than 100 assignments. I led a distributed, round-the-clock emergency "
    "response team that delivered a complete containment-riser design in eight weeks by "
    "repurposing existing assets. That work required clear technical decisions and coordinated "
    "execution under pressure."
)
DIGITAL = (
    "At Occidental, I coordinated engineering, data-science and software teams delivering "
    "more than 50 physics-based algorithms supporting approximately 20,000 wells. I owned "
    "requirements, validation and production support. In marine analysis, I have built "
    "Python/API workflows for model preparation, execution, results extraction and reporting."
)
ROLES = {
    "AWS": {
        "title": "Sr. Mechanical Engineer - CFD, Global Engineering Strategy",
        "subtitle": "Mechanical Engineering and CFD",
        "profile": "Seeking to contribute offshore industrial engineering, simulation and technical review experience to AWS data-center engineering. Direct data-center cooling design experience is not claimed.",
        "intro": "I am applying for the Sr. Mechanical Engineer - CFD position in Global Engineering Strategy (job 10437257). I bring 23 years of offshore and marine engineering experience, a Texas PE license, and recent CFD work. The role's combination of industrial engineering, CFD scope definition and technical review provides a practical route for applying that experience to data-center infrastructure.",
        "middle": [CFD, DELIVERY],
        "close": "I would bring simulation judgment, documented engineering decisions and experience turning analysis into deliverable recommendations. My background is offshore engineering; I would need to develop AWS-specific site cooling, water-efficiency and data-center CFD expertise. I hold an MS in Mechanical Engineering from Texas A&M and a B.Tech. in Mechanical Engineering from IIT Madras, and would welcome a discussion about the fit.",
    },
    "Armada": {
        "title": "Senior Mechanical Design Engineer – Liquid Cooling Systems",
        "subtitle": "Mechanical Analysis and Engineering Delivery",
        "profile": "Seeking to bring mechanical analysis, interface-risk assessment and multidisciplinary delivery experience to Armada's modular infrastructure team. This is a transition into liquid-cooling systems, not a claim of prior data-center cooling design.",
        "intro": "I am applying for the Senior Mechanical Design Engineer – Liquid Cooling Systems role. I bring 23 years of offshore engineering and a Texas PE license, with experience combining mechanical analysis, design reviews and multidisciplinary delivery. Armada's emphasis on repeatable modular infrastructure interests me as an application of practical engineering judgment and disciplined interfaces.",
        "middle": [DELIVERY, CFD],
        "close": "I would bring experience documenting design assumptions, evaluating failure mechanisms and building repeatable analysis workflows. My work on offshore pipelines, risers and installation analysis does not establish cooling-loop hydraulic design, pump selection or CDU commissioning experience. I would need to build that specific expertise. I hold mechanical engineering degrees from Texas A&M and IIT Madras and would welcome a discussion about whether my analysis and delivery background can contribute to your team.",
    },
    "OpenAI": {
        "title": "Mechanical Commissioning Lead",
        "subtitle": "Mechanical Engineering Assurance and Delivery",
        "profile": "Seeking to contribute design-review discipline, analysis validation and multidisciplinary engineering leadership to OpenAI's infrastructure team. This is a transition toward commissioning; direct data-center commissioning, FAT/SAT execution and controls verification experience are not claimed.",
        "intro": "I am applying for the Mechanical Commissioning Lead role. I bring 23 years of marine and offshore engineering, a Texas PE license, and experience in design reviews, engineering risk assessment and multidisciplinary delivery. I am interested in applying that judgment to the reliable delivery of OpenAI's physical infrastructure.",
        "middle": [DELIVERY, DIGITAL],
        "close": "My recent CFD work and academic thesis work in fluid mechanics and heat transfer provide a relevant analytical foundation. I recognize that this role requires hands-on cooling and hydronic commissioning, FAT/SAT and controls verification that my existing experience does not establish. My strength is engineering assurance and coordinated technical delivery; this would be a deliberate transition into data-center commissioning. I hold mechanical engineering degrees from Texas A&M and IIT Madras and would welcome an assessment of whether that foundation fits your needs.",
    },
}


def replace_text(paragraph, text):
    """Preserve paragraph styling and first-run typography."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def make_cv(employer, role):
    doc = Document(SOURCE / "Vamsee_Achanta_Anthropic_CV.docx")
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Subtitle":
            replace_text(paragraph, role["subtitle"])
        elif paragraph.text.startswith("Texas-licensed Professional Engineer with"):
            replace_text(paragraph, COMMON + role["profile"])
    doc.core_properties.title = f"Vamsee Achanta {employer} CV"
    doc.save(ROOT / f"Vamsee_Achanta_{employer}_CV.docx")


def make_letter(employer, role):
    doc = Document(SOURCE / "Vamsee_Achanta_Anthropic_Cover_Letter.docx")
    original = doc.paragraphs
    texts = ["4 October 2026", f"Dear {employer} Hiring Team,", role["intro"],
             *role["middle"], role["close"], "Sincerely,\nVamsee Achanta, P.E."]
    replace_text(original[3], "Application for " + role["title"])
    for paragraph, text in zip(original[4:], texts):
        replace_text(paragraph, text)
    for paragraph in original[4 + len(texts):]:
        paragraph._element.getparent().remove(paragraph._element)
    doc.core_properties.title = f"Vamsee Achanta {employer} Cover Letter"
    doc.save(ROOT / f"Vamsee_Achanta_{employer}_Cover_Letter.docx")


def validate():
    for path in ROOT.glob("*.docx"):
        text = "\n".join(p.text for p in Document(path).paragraphs)
        assert "Anthropic" not in text, path
        assert "Vamsee Achanta" in text, path
        if path.stem.endswith("_CV"):
            assert "Dec 2023 – Jun 2026" in text, path
            assert "Jun 2015 – Nov 2023" in text, path
        print(path.name, len(text.split()), "words; checked")


if __name__ == "__main__":
    for employer, role in ROLES.items():
        make_cv(employer, role)
        make_letter(employer, role)
    validate()
