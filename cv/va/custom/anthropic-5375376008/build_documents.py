"""Build the private Anthropic CV and cover letter from candidate evidence."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.65)
    sec.left_margin = sec.right_margin = Inches(0.75)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.03
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Title'].paragraph_format.space_after = Pt(3)
    doc.styles['Heading 1'].font.size = Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.core_properties.author = 'Vamsee Achanta'
    return doc


def header(doc, subject):
    doc.add_paragraph('Vamsee Achanta', 'Title')
    doc.add_paragraph('Texas Professional Engineer | Houston, Texas')
    doc.add_paragraph('achantav@gmail.com | 713-306-9029 | linkedin.com/in/vamseeachanta')
    doc.add_paragraph(subject, 'Subtitle')


def make_cv():
    doc = new_document()
    header(doc, 'Mechanical Engineering and CFD')
    lines = (ROOT / 'resume-draft.txt').read_text(encoding='utf-8').splitlines()[5:]
    for line in lines:
        if not line:
            continue
        if line == 'Engineering SME, Data Science Team — Occidental Petroleum, Houston, TX':
            heading = doc.add_paragraph('Professional Experience Continued', 'Heading 1')
            heading.paragraph_format.page_break_before = True
        if line.isupper():
            title = line.title().replace('Cfd', 'CFD').replace('Fluid-Mechanics', 'Fluid Mechanics')
            doc.add_paragraph(title, 'Heading 1')
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], 'List Bullet')
        else:
            p = doc.add_paragraph(line)
            if ' — ' in line and ('Houston' in line or 'Inc.' in line):
                p.runs[0].bold = True
                p.paragraph_format.keep_with_next = True
            if line.startswith(('Dec 2023', 'Jun 2016', 'Sep 2017', 'Jun 2015', 'Aug 2003')):
                p.paragraph_format.keep_with_next = True
    doc.core_properties.title = 'Vamsee Achanta Anthropic CV'
    doc.save(ROOT / 'Vamsee_Achanta_Anthropic_CV.docx')


def make_letter():
    doc = new_document()
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(11)
    header(doc, 'Application for Data Center Mechanical Engineer')
    paragraphs = [
        '4 October 2026',
        'Dear Anthropic Hiring Team,',
        'I am applying for the Data Center Mechanical Engineer role. I bring 23 years of marine and offshore engineering experience, recent computational fluid dynamics modelling and analysis, and academic thesis work in fluid mechanics and heat transfer. I am keen to apply that foundation to the mechanical infrastructure supporting Anthropic’s AI systems. This would be a deliberate transition from offshore engineering into data-center engineering.',
        'My recent CFD work includes OpenFOAM modelling of laminar, turbulent and free-surface flows, with benchmark-focused cases involving boundary layers, vortex shedding, airfoil flow and dam-break behaviour. My engineering work also includes COMSOL multiphysics modelling combining fluid mechanics, mass transfer, diffusion and electrochemistry. Together with my thesis work, these experiences provide a foundation for reasoning about flow and transport while developing the cooling-system expertise this role requires.',
        'I have applied engineering analysis under demanding delivery conditions. At 2H Offshore, I led a distributed, round-the-clock team that delivered a containment-riser design in eight weeks by repurposing existing assets. Across more than 100 assignments, my responsibilities included design reviews, technical reports, client communication and multidisciplinary delivery. At Occidental, I coordinated engineering, data-science and software teams delivering more than 50 physics-based algorithms supporting approximately 20,000 wells. I have also built repeatable workflows for model preparation, execution and reporting.',
        'Anthropic’s need to connect mechanical design decisions with reliable infrastructure delivery interests me. I would bring simulation experience, design-review discipline and practical engineering judgment. I recognize that data-center cooling architectures, building-services codes and facility commissioning require domain-specific experience that I would need to develop; I would not present my offshore work as prior data-center delivery.',
        'I hold a master’s degree in Mechanical Engineering from Texas A&M University, a bachelor’s degree in Mechanical Engineering from IIT Madras, and a Texas Professional Engineer license. I would welcome a discussion about how my engineering work and fluid-mechanics foundation could contribute to your team.',
        'Sincerely,\nVamsee Achanta, P.E.',
    ]
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.core_properties.title = 'Vamsee Achanta Anthropic Cover Letter'
    doc.save(ROOT / 'Vamsee_Achanta_Anthropic_Cover_Letter.docx')


if __name__ == '__main__':
    make_cv()
    make_letter()
